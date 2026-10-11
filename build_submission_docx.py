"""Build submission PDF:
1) Use Uni/title page.docx PAGE 1 exactly as-is (logo/layout untouched)
2) Append a properly formatted technical report
"""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parent
TITLE_SRC = Path(r"e:\Code\Uni\title page.docx")
RESULTS = ROOT / "results"
OUT_DOCX = ROOT / "report" / "Technical_Report.docx"
OUT_PDF = ROOT / "report" / "Technical_Report.pdf"
BODY_DOCX = ROOT / "report" / "_report_body.docx"

GITHUB = "https://github.com/azhar-fullstack/RL-Assignment-CliffWalking-QLearning"
NAME = "Azhar Mehmood"
REG = "FA23-BAI-039"


def set_run_font(run, size=11, bold=False, italic=False, color=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=6, line=1.1):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
        set_paragraph_spacing(p, before=10, after=4, line=1.1)
    else:
        set_run_font(run, size=11.5, bold=True, color=RGBColor(0x2F, 0x54, 0x96))
        set_paragraph_spacing(p, before=8, after=3, line=1.1)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    set_paragraph_spacing(p, before=0, after=6, line=1.1)
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True)
    set_paragraph_spacing(p, before=4, after=6, line=1.1)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=RGBColor(0x40, 0x40, 0x40))
    set_paragraph_spacing(p, before=1, after=8, line=1.0)


def add_picture(doc, path: Path, width_in=6.0):
    if not path.exists():
        add_body(doc, f"[Missing figure: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=1, line=1.0)
    p.add_run().add_picture(str(path), width=Inches(width_in))


def set_cell_shading(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_hyperparam_table(doc: Document):
    headers = ["Hyperparameter", "Symbol", "Value", "Why I chose it"]
    rows = [
        ["Learning rate", "alpha", "0.5", "0.1 was slow; 0.9 less stable"],
        ["Discount factor", "gamma", "0.99", "Care more about reaching goal"],
        ["Start epsilon", "eps0", "1.0", "Start fully random"],
        ["Min epsilon", "eps_min", "0.01", "Keep a little exploration"],
        ["Epsilon decay", "eps_decay", "0.995", "Decay once per episode"],
        ["Training episodes", "N", "5000", "Main training run"],
        ["Max steps / episode", "T", "200", "Stop very long episodes"],
        ["Eval episodes", "-", "100", "Greedy policy test"],
        ["Seed", "-", "42", "Fair comparison"],
    ]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.autofit = True
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "666666")
        borders.append(elem)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(cell, "4472C4")

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=9)
            if i % 2 == 0:
                set_cell_shading(cell, "F2F2F2")
    doc.add_paragraph()


def build_report_body_docx() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Technical Report")
    set_run_font(run, size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    set_paragraph_spacing(p, before=0, after=3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Custom Reinforcement Learning Agent Implementation using Farama Gymnasium"
    )
    set_run_font(run, size=11.5, bold=True)
    set_paragraph_spacing(p, before=0, after=2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f"{NAME}  |  {REG}  |  CliffWalking-v1  |  Q-Learning (from scratch)"
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x40, 0x40, 0x40))
    set_paragraph_spacing(p, before=0, after=10)

    add_heading(doc, "1. Problem Description", 1)
    add_body(
        doc,
        "For this assignment I picked CliffWalking-v1 from Gymnasium. It is a small 4x12 grid "
        "(48 states), so I could implement the algorithm myself without neural networks. The agent "
        "starts at the bottom-left and must reach the goal at the bottom-right. A cliff sits along "
        "the bottom row; stepping into it gives a big penalty and sends the agent back to the start. "
        "Actions are: 0 = Up, 1 = Right, 2 = Down, 3 = Left.",
    )
    add_body(
        doc,
        "Every normal step gives reward -1. Falling into the cliff gives -100. Reaching the goal "
        "ends the episode. The best safe path is 13 steps, so the best score is -13. A longer upper "
        "path is safer in some sense but gets a worse (more negative) score. That tradeoff is why "
        "this environment is useful for testing RL.",
    )

    add_heading(doc, "2. Methodology", 1)
    add_heading(doc, "2.1 Algorithm", 2)
    add_body(
        doc,
        "I implemented Q-Learning with an epsilon-greedy policy. I keep a table Q(s, a) for every "
        "state-action pair. Most of the time the agent picks the best action from the table, but "
        "with probability epsilon it picks a random action to explore. After each step:",
    )
    add_formula(doc, "Q(s,a) = Q(s,a) + alpha * [ r + gamma * max Q(s', a') - Q(s,a) ]")
    add_body(
        doc,
        "Here alpha is the learning rate and gamma is the discount factor. Epsilon starts at 1.0 "
        "and decays each episode down to 0.01. Using max over next actions makes this the normal "
        "off-policy Q-Learning update.",
    )
    add_heading(doc, "2.2 Code Structure", 2)
    add_body(
        doc,
        "I split the project into three files. agent.py has the QLearningAgent class (action "
        "selection, Q update, epsilon decay, save/load). utils.py has training, evaluation, "
        "plots, and GIF rendering. main.py runs the full pipeline. The environment is created with "
        "gymnasium.make(\"CliffWalking-v1\"). After training I save models/q_table.npy and "
        "results/agent_demo.gif.",
    )

    add_heading(doc, "3. Experimental Setup and Hyperparameters", 1)
    add_body(
        doc,
        "I trained in Python with Gymnasium. The table below shows the final settings used for the "
        "saved model and GIF. For sensitivity tests I ran 3000 episodes again and changed one "
        "setting at a time with the same seed.",
    )
    add_hyperparam_table(doc)

    add_heading(doc, "4. Results and Discussion", 1)
    add_body(
        doc,
        "After 5000 episodes, greedy evaluation over 100 episodes gave average return -13.0 "
        "(no variance) and 100% success. So the agent learned the optimal 13-step path. In "
        "Figure 1, reward improves mainly in the first 1000-2000 episodes, then becomes more stable "
        "as epsilon gets small. Figure 2 shows higher values along the good path and lower values "
        "near the cliff.",
    )
    add_picture(doc, RESULTS / "training_curves.png", width_in=6.1)
    add_caption(doc, "Figure 1. Reward, success rate, and epsilon during training.")
    add_picture(doc, RESULTS / "q_value_heatmap.png", width_in=5.8)
    add_caption(doc, "Figure 2. Max Q-value for each cell on the 4x12 map.")
    add_body(
        doc,
        "In the hyperparameter check (Figure 3), alpha = 0.5 beat alpha = 0.1 under the same "
        "episode budget. Gamma = 0.99 worked fine. Slow epsilon decay (0.999) kept exploration "
        "high for longer, so late training return stayed worse. Learning rate and epsilon decay "
        "mattered the most in my runs.",
    )
    add_picture(doc, RESULTS / "hyperparameter_sensitivity.png", width_in=5.5)
    add_caption(doc, "Figure 3. Average return in the last 500 episodes for different settings.")

    add_heading(doc, "4.1 Problems I Noticed", 2)
    add_body(
        doc,
        "Early on the agent falls a lot because epsilon is high; that looks bad in the reward "
        "curve but is normal while exploring. It can also learn a longer upper-row path that is "
        "safer but worse than -13. This table method works here, but not for continuous tasks "
        "like CartPole without something like DQN. The demo GIF shows the final greedy policy "
        "reaching the goal without falling.",
    )

    add_heading(doc, "5. Conclusion", 1)
    add_body(
        doc,
        "I built a from-scratch Q-Learning agent on CliffWalking-v1, trained it, plotted the "
        "curves, saved the Q-table, and recorded a short GIF. In evaluation it reached the "
        "optimal score (-13). Changing alpha, gamma, and epsilon decay clearly changed learning "
        "speed and final return. Next I would try SARSA or Double Q-Learning.",
    )

    p = doc.add_paragraph()
    run = p.add_run("GitHub: ")
    set_run_font(run, size=10.5, bold=True)
    run = p.add_run(GITHUB)
    set_run_font(run, size=10, color=RGBColor(0x05, 0x63, 0xC1))
    set_paragraph_spacing(p, before=4, after=3)

    p = doc.add_paragraph()
    run = p.add_run(
        "Files: main.py, agent.py, utils.py, notebooks/training.ipynb, "
        "models/q_table.npy, results/agent_demo.gif"
    )
    set_run_font(run, size=9.5)
    set_paragraph_spacing(p, before=0, after=2)

    BODY_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(BODY_DOCX))
    print(f"Wrote body DOCX: {BODY_DOCX}")
    return BODY_DOCX


def word_export_pdf(docx_path: Path, pdf_path: Path) -> None:
    try:
        import win32com.client

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
            doc.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
            doc.Close(False)
        finally:
            word.Quit()
    except Exception:
        ps = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$doc = $word.Documents.Open('{(str(docx_path.resolve())).replace("'", "''")}')
$doc.ExportAsFixedFormat('{(str(pdf_path.resolve())).replace("'", "''")}', 17)
$doc.Close($false)
$word.Quit()
"""
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print(f"Exported PDF: {pdf_path}")


def build_combined_docx_for_editing(title_pdf_page1_note: str = "") -> None:
    """Also keep a DOCX copy: title template + body, for students who need .docx."""
    # For DOCX submission convenience: body only is not enough.
    # We keep Technical_Report.docx as body with a note that PDF has official cover.
    # Better: produce final PDF only as required.
    pass


def main() -> None:
    if not TITLE_SRC.exists():
        raise FileNotFoundError(TITLE_SRC)

    OUT_PDF.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        title_pdf = tmp_dir / "title.pdf"
        body_pdf = tmp_dir / "body.pdf"

        # 1) Export original title page exactly via Word
        word_export_pdf(TITLE_SRC, title_pdf)

        # 2) Build and export report body
        body_docx = build_report_body_docx()
        word_export_pdf(body_docx, body_pdf)

        # 3) Merge: ONLY first page of title template + full report body
        writer = PdfWriter()
        title_reader = PdfReader(str(title_pdf))
        body_reader = PdfReader(str(body_pdf))
        writer.add_page(title_reader.pages[0])  # as-is first page only
        for page in body_reader.pages:
            writer.add_page(page)

        with OUT_PDF.open("wb") as f:
            writer.write(f)

        # Save a convenient DOCX: copy body as Technical_Report.docx and also
        # keep a note file? User asked for PDF mainly. Still save body docx path
        # as Technical_Report.docx for editing if needed.
        OUT_DOCX.write_bytes(body_docx.read_bytes())

    reader = PdfReader(str(OUT_PDF))
    print(f"Wrote final PDF: {OUT_PDF} ({len(reader.pages)} pages)")
    print("Page 1 = original Uni/title page.docx first page (logo preserved).")


if __name__ == "__main__":
    main()
